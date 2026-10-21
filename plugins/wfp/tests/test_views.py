import io

from datetime import date, datetime, timezone

import openpyxl

from django.urls import reverse

from iaso import models as m
from iaso.test import TestCase
from plugins.wfp.models import Beneficiary, Journey, Step, Visit


class BeneficiariesViewTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.account = m.Account.objects.create(name="WFP")
        cls.other_account = m.Account.objects.create(name="Other")
        cls.form = m.Form.objects.create(name="Reference", form_id="reference")
        cls.entity_type = m.EntityType.objects.create(name="Child", reference_form=cls.form, account=cls.account)
        cls.other_entity_type = m.EntityType.objects.create(
            name="Other child", reference_form=cls.form, account=cls.other_account
        )
        cls.user = cls.create_user_with_profile(username="user", account=cls.account)

        cls.reference = m.Instance.objects.create(form=cls.form, json={"first_name": "Amina", "household": {"size": 5}})
        cls.entity = m.Entity.objects.create(entity_type=cls.entity_type, account=cls.account, attributes=cls.reference)
        cls.reference.entity = cls.entity
        cls.reference.save()
        cls.visit_instance = m.Instance.objects.create(
            form=cls.form, entity=cls.entity, json={"muac": "13.1", "referrer_name": "Fatou"}
        )
        cls.beneficiary = Beneficiary.objects.create(
            entity=cls.entity, account=cls.account, birth_date=date(2024, 1, 1), gender="Female"
        )
        cls.journey = Journey.objects.create(
            beneficiary=cls.beneficiary, nutrition_programme="BSFP", programme_type="U5", start_date=date(2026, 5, 1)
        )
        cls.visit = Visit.objects.create(
            journey=cls.journey,
            date=datetime(2026, 5, 1, 12, tzinfo=timezone.utc),
            muac_size="13.1",
            instance_id=cls.visit_instance.id,
        )
        Step.objects.create(
            visit=cls.visit, assistance_type="lns", quantity_given=14, instance_id=cls.visit_instance.id
        )

        cls.other_entity = m.Entity.objects.create(entity_type=cls.other_entity_type, account=cls.other_account)
        cls.other_beneficiary = Beneficiary.objects.create(
            entity=cls.other_entity, account=cls.other_account, birth_date=date(2023, 1, 1), gender="Male"
        )

    def test_requires_login(self):
        response = self.client.get(reverse("wfp_beneficiaries"))
        self.assertEqual(response.status_code, 302)

    def test_list_is_limited_to_the_user_account(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("wfp_beneficiaries"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual([b.id for b in response.context["page"]], [self.beneficiary.id])
        self.assertContains(response, str(self.entity.uuid))
        self.assertNotContains(response, str(self.other_entity.uuid))

    def test_list_shows_visits_and_counts(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("wfp_beneficiaries"))
        beneficiary = response.context["page"][0]
        self.assertEqual(beneficiary.journeys_count, 1)
        self.assertEqual(beneficiary.visits_count, 1)
        self.assertEqual(list(beneficiary.journeys[0].visit_set.all()), [self.visit])
        self.assertContains(response, "lns: 14.0")

    def test_programme_filter_keeps_all_journeys_in_counts(self):
        Journey.objects.create(beneficiary=self.beneficiary, nutrition_programme="TSFP", programme_type="U5")
        self.client.force_login(self.user)
        response = self.client.get(reverse("wfp_beneficiaries"), {"programme": "BSFP"})
        self.assertEqual([b.journeys_count for b in response.context["page"]], [2])
        response = self.client.get(reverse("wfp_beneficiaries"), {"programme": "OTP"})
        self.assertEqual(list(response.context["page"]), [])

    def test_search_by_entity_id(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("wfp_beneficiaries"), {"q": str(self.entity.id)})
        self.assertEqual([b.id for b in response.context["page"]], [self.beneficiary.id])

    def test_beneficiary_data_shows_reference_and_submissions(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("wfp_beneficiary_data", args=[self.entity.id]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Amina")
        self.assertContains(response, "household/size")
        self.assertContains(response, "Fatou")
        self.assertEqual({s.id for s in response.context["submissions"]}, {self.reference.id, self.visit_instance.id})

    def test_beneficiary_data_of_another_account_is_not_found(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("wfp_beneficiary_data", args=[self.other_entity.id]))
        self.assertEqual(response.status_code, 404)

    def test_submission_fields(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("wfp_submission_fields", args=[self.visit_instance.id]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "referrer_name")
        self.assertContains(response, "Fatou")

    def test_submission_fields_of_another_account_is_not_found(self):
        other_instance = m.Instance.objects.create(form=self.form, entity=self.other_entity, json={"secret": "x"})
        self.client.force_login(self.user)
        response = self.client.get(reverse("wfp_submission_fields", args=[other_instance.id]))
        self.assertEqual(response.status_code, 404)

    def test_list_links_submissions_to_their_fields(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("wfp_beneficiaries"))
        self.assertContains(response, reverse("wfp_submission_fields", args=[self.visit_instance.id]))


class ExportTestData(TestCase):
    """An entity with a reference submission, a visit and a medical submission, plus one in another account."""

    @classmethod
    def setUpTestData(cls):
        cls.account = m.Account.objects.create(name="WFP")
        cls.other_account = m.Account.objects.create(name="Other")
        cls.user = cls.create_user_with_profile(username="user", account=cls.account)
        cls.reference_form = m.Form.objects.create(name="Reference", form_id="reference")
        cls.visit_form = m.Form.objects.create(name="Visit", form_id="visit")
        # the medical form has a PII configuration: only `passing_urine` is not personal data
        cls.medical_form = m.Form.objects.create(name="Medical", form_id="medical", non_pii_fields=["passing_urine"])
        entity_type = m.EntityType.objects.create(name="Child", reference_form=cls.reference_form, account=cls.account)

        country = m.OrgUnit.objects.create(name="South Sudan")
        state = m.OrgUnit.objects.create(name="Unity", parent=country)
        cls.facility = m.OrgUnit.objects.create(name="Bentiu PHCC", parent=state)

        cls.entity = m.Entity.objects.create(entity_type=entity_type, account=cls.account)
        day = datetime(2026, 5, 1, 9, tzinfo=timezone.utc)
        reference = m.Instance.objects.create(
            form=cls.reference_form,
            entity=cls.entity,
            org_unit=cls.facility,
            source_created_at=day,
            json={"first_name": "Amina", "is_twin": "no"},
        )
        cls.entity.attributes = reference
        cls.entity.save()
        anchor = m.Instance.objects.create(
            form=cls.visit_form,
            entity=cls.entity,
            org_unit=cls.facility,
            source_created_at=day.replace(hour=10),
            json={"muac": "13.1", "weight_kgs": "9.5"},
        )
        m.Instance.objects.create(
            form=cls.medical_form,
            entity=cls.entity,
            org_unit=cls.facility,
            source_created_at=day.replace(hour=11),
            json={"Passing_Urine": "yes", "fever": "1"},
        )
        beneficiary = Beneficiary.objects.create(
            entity=cls.entity, account=cls.account, birth_date=date(2024, 1, 15), gender="Female", guidelines="NEW"
        )
        journey = Journey.objects.create(
            beneficiary=beneficiary,
            nutrition_programme="TSFP",
            programme_type="U5",
            admission_criteria="muac",
            start_date=date(2026, 5, 1),
            instance_id=anchor.id,
        )
        visit = Visit.objects.create(
            journey=journey, date=day, number=1, org_unit=cls.facility, muac_size="13.1", instance_id=anchor.id
        )
        Step.objects.create(visit=visit, assistance_type="Soap", quantity_given=1, instance_id=anchor.id)

        other_entity_type = m.EntityType.objects.create(
            name="Other", reference_form=cls.reference_form, account=cls.other_account
        )
        cls.other_entity = m.Entity.objects.create(entity_type=other_entity_type, account=cls.other_account)
        Beneficiary.objects.create(entity=cls.other_entity, account=cls.other_account, birth_date=date(2023, 1, 1))


class CodaExportTestCase(ExportTestData):
    def _export(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("wfp_beneficiaries_coda_export"))
        self.assertEqual(response.status_code, 200)
        workbook = openpyxl.load_workbook(io.BytesIO(b"".join(response.streaming_content)))
        rows = list(workbook.active.iter_rows(values_only=True))
        headers = rows[1]
        return [dict(zip(headers, row)) for row in rows[2:]], rows[0]

    def test_one_row_per_visit_of_the_account(self):
        rows, steps = self._export()
        self.assertEqual(steps[0], "Beneficiary details")
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row["beneficiary unique Identification"], str(self.entity.uuid))
        self.assertEqual(row["gender"], "Female")
        self.assertEqual(row["age (months for children, years for women)"], 27)
        self.assertEqual(row["health center name"], "Bentiu PHCC")
        self.assertEqual(row["state"], "Unity")
        self.assertEqual(row["State"], "Unity")
        self.assertEqual(row["programme"], "TSFP")
        self.assertEqual(row["treatment protocol"], "NEW")
        self.assertEqual(row["Is  registration visit"], "yes")
        self.assertEqual(row["MUAC"], "13.1")
        self.assertEqual(row["weight"], "9.5")
        self.assertEqual(row["soap"], "yes")
        self.assertEqual(row["Is twin?"], "no")

    def test_answers_of_the_visit_submissions_are_matched_loosely(self):
        rows, _ = self._export()
        self.assertEqual(rows[0]["passing Urine?"], "yes")

    def test_personal_data_is_not_exported(self):
        rows, _ = self._export()
        # `fever` is not whitelisted in the PII configuration of the medical form
        self.assertIsNone(rows[0]["fever"])

    def test_export_uses_the_filters_of_the_screen(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("wfp_beneficiaries_coda_export"), {"programme": "OTP"})
        workbook = openpyxl.load_workbook(io.BytesIO(b"".join(response.streaming_content)))
        self.assertEqual(len(list(workbook.active.iter_rows())), 2)  # the two header rows only


class RawExportTestCase(ExportTestData):
    def _export(self, params=None):
        self.client.force_login(self.user)
        response = self.client.get(reverse("wfp_beneficiaries_raw_export"), params or {})
        self.assertEqual(response.status_code, 200)
        workbook = openpyxl.load_workbook(io.BytesIO(b"".join(response.streaming_content)))
        sheets = {}
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            sheets[sheet.title] = [dict(zip(rows[0], row)) for row in rows[1:]] if rows else []
        return sheets

    def test_sheets(self):
        sheets = self._export()
        self.assertEqual(
            list(sheets),
            ["README", "beneficiaries", "journeys", "visits", "steps", "medical", "reference", "visit", "forms"],
        )

    def test_etl_tables_of_the_account(self):
        sheets = self._export()
        self.assertEqual([row["entity_uuid"] for row in sheets["beneficiaries"]], [str(self.entity.uuid)])
        self.assertEqual(len(sheets["journeys"]), 1)
        self.assertEqual(len(sheets["visits"]), 1)
        self.assertEqual(sheets["steps"][0]["assistance_type"], "Soap")

    def test_submissions_are_tied_to_the_etl_rows(self):
        sheets = self._export()
        visit = sheets["visits"][0]
        submission = sheets["visit"][0]
        self.assertEqual(submission["instance_id"], visit["instance_id"])
        self.assertEqual(submission["beneficiary_id"], visit["beneficiary_id"])
        self.assertEqual(submission["entity_uuid"], str(self.entity.uuid))
        self.assertEqual(submission["muac"], "13.1")
        self.assertEqual(sheets["reference"][0]["first_name"], "Amina")

    def test_personal_data_is_not_exported(self):
        medical = self._export()["medical"]
        self.assertEqual(medical[0]["Passing_Urine"], "yes")
        self.assertNotIn("fever", medical[0])

    def test_export_uses_the_filters_of_the_screen(self):
        sheets = self._export({"programme": "OTP"})
        self.assertEqual(sheets["beneficiaries"], [])
        self.assertEqual(sheets["forms"], [])
